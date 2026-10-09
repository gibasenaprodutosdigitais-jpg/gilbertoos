"""
Formata um capitulo do livro nas normas que a editora passou.

Times New Roman 12, espacamento 1,5, margens 3cm (superior e esquerda) e
2cm (inferior e direita), alinhamento justificado, recuo de paragrafo de
1,25cm e numero de pagina no canto superior direito.

Gera dois arquivos: o .docx, que e o que vai para a editora, e um .html com
as mesmas medidas, que serve para contar paginas e olhar antes de mandar.

Uso:  python3 formatar-capitulo.py capitulos/07-do-zero-ao-grupo-sena.md
"""
import io
import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

AQUI = Path(__file__).parent
SAIDA = AQUI / "formatado"

# ---------------------------------------------------------------- leitura

def limpar(md: str) -> str:
    """Tira o que e nota de trabalho e nao faz parte do capitulo."""
    # o bloco de citacao do topo (as notas de versao)
    md = re.sub(r"\A(# [^\n]+\n\n)(> [^\n]*\n)+\n?", r"\1", md)
    # tudo a partir das notas de redacao
    md = re.split(r"\n---\n\n## Notas de redação", md)[0]
    return md.rstrip().rstrip("-").rstrip() + "\n"


def partir(md: str):
    """Quebra o markdown em blocos (tipo, conteudo)."""
    blocos, buf, dentro_tabela = [], [], []
    def fecha_paragrafo():
        if buf:
            blocos.append(("p", " ".join(buf).strip()))
            buf.clear()
    def fecha_tabela():
        if dentro_tabela:
            blocos.append(("tabela", list(dentro_tabela)))
            dentro_tabela.clear()

    for linha in md.split("\n"):
        crua = linha.rstrip()
        if crua.startswith("|"):
            fecha_paragrafo()
            dentro_tabela.append(crua)
            continue
        fecha_tabela()
        if not crua.strip():
            fecha_paragrafo()
        elif crua.startswith("### "):
            fecha_paragrafo(); blocos.append(("h3", crua[4:]))
        elif crua.startswith("## "):
            fecha_paragrafo(); blocos.append(("h2", crua[3:]))
        elif crua.startswith("# "):
            fecha_paragrafo(); blocos.append(("h1", crua[2:]))
        elif crua.startswith("> "):
            fecha_paragrafo()
            linha_cita = crua[2:]
            titulo = re.fullmatch(r"\*\*[^*]+\*\*", linha_cita.strip())
            # linhas seguidas de citacao formam um paragrafo so; o titulo em
            # negrito da Sacada fica sozinho, como cabecalho do box
            if blocos and blocos[-1][0] == "cita" and not titulo \
               and not re.fullmatch(r"\*\*[^*]+\*\*", blocos[-1][1].strip()):
                blocos[-1] = ("cita", blocos[-1][1] + " " + linha_cita)
            else:
                blocos.append(("cita", linha_cita))
        elif crua.strip() == "---":
            fecha_paragrafo(); blocos.append(("regra", ""))
        elif re.match(r"^\d+\. ", crua):
            fecha_paragrafo(); blocos.append(("num", re.sub(r"^\d+\. ", "", crua)))
        elif crua.startswith("- "):
            fecha_paragrafo(); blocos.append(("item", crua[2:]))
        elif crua.startswith("    ") and blocos and blocos[-1][0] in ("num", "item"):
            tipo, texto = blocos.pop()
            blocos.append((tipo, texto + " " + crua.strip()))
        else:
            buf.append(crua.strip())
    fecha_paragrafo(); fecha_tabela()
    return blocos


# ------------------------------------------------------------------ docx

def escrever_inline(par, texto, italico_base=False):
    """Aplica **negrito**, *italico* e `codigo` dentro de um paragrafo."""
    for pedaco in re.split(r"(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`)", texto):
        if not pedaco:
            continue
        r = par.add_run()
        if pedaco.startswith("**") and pedaco.endswith("**"):
            r.text, r.bold = pedaco[2:-2], True
        elif pedaco.startswith("*") and pedaco.endswith("*"):
            r.text, r.italic = pedaco[1:-1], True
        elif pedaco.startswith("`") and pedaco.endswith("`"):
            r.text = pedaco[1:-1]
        else:
            r.text = pedaco
        if italico_base:
            r.italic = True


def corpo(doc, texto, recuo=True, italico=False):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    pf.line_spacing = 1.5
    pf.space_after = Pt(0)
    pf.first_line_indent = Cm(1.25) if recuo else Cm(0)
    escrever_inline(p, texto, italico)
    return p


def numero_de_pagina(secao):
    """Campo PAGE no cabecalho, alinhado a direita."""
    p = secao.header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = p.add_run()
    for tipo, valor in (("begin", None), ("instrText", "PAGE"), ("end", None)):
        el = OxmlElement("w:fldChar" if tipo != "instrText" else "w:instrText")
        if tipo == "instrText":
            el.set(qn("xml:space"), "preserve"); el.text = " PAGE "
        else:
            el.set(qn("w:fldCharType"), tipo)
        r._r.append(el)
    r.font.name = "Times New Roman"
    r.font.size = Pt(12)


def montar_docx(blocos, destino: Path, titulo: str):
    doc = Document()
    est = doc.styles["Normal"]
    est.font.name = "Times New Roman"
    est.font.size = Pt(12)
    est.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")

    s = doc.sections[0]
    s.top_margin, s.left_margin = Cm(3), Cm(3)
    s.bottom_margin, s.right_margin = Cm(2), Cm(2)
    s.header_distance = Cm(1.5)
    numero_de_pagina(s)

    for tipo, conteudo in blocos:
        if tipo == "h1":
            p = doc.add_paragraph()
            p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.line_spacing = 1.5
            p.paragraph_format.space_after = Pt(18)
            r = p.add_run(conteudo.upper()); r.bold = True; r.font.size = Pt(14)
        elif tipo in ("h2", "h3"):
            p = doc.add_paragraph()
            p.paragraph_format.line_spacing = 1.5
            p.paragraph_format.space_before = Pt(18)
            p.paragraph_format.space_after = Pt(6)
            r = p.add_run(conteudo); r.bold = True
            r.italic = (tipo == "h3")
        elif tipo == "p":
            corpo(doc, conteudo)
        elif tipo == "cita":
            p = corpo(doc, conteudo, recuo=False, italico=True)
            p.paragraph_format.left_indent = Cm(2)
            p.paragraph_format.space_before = Pt(6)
        elif tipo == "num":
            p = corpo(doc, conteudo, recuo=False)
            p.paragraph_format.left_indent = Cm(1.25)
            p.paragraph_format.space_before = Pt(6)
        elif tipo == "item":
            p = corpo(doc, "• " + conteudo, recuo=False)
            p.paragraph_format.left_indent = Cm(1.25)
        elif tipo == "regra":
            pass
        elif tipo == "tabela":
            linhas = [l for l in conteudo if not re.match(r"^\|[\s|:-]+\|$", l)]
            celulas = [[c.strip() for c in l.strip("|").split("|")] for l in linhas]
            if not celulas:
                continue
            t = doc.add_table(rows=len(celulas), cols=len(celulas[0]))
            t.style = "Table Grid"
            t.alignment = WD_TABLE_ALIGNMENT.CENTER
            for i, linha in enumerate(celulas):
                for j, cel in enumerate(linha):
                    cp = t.cell(i, j).paragraphs[0]
                    cp.paragraph_format.line_spacing = 1.0
                    cp.paragraph_format.space_after = Pt(0)
                    escrever_inline(cp, cel)
                    for r in cp.runs:
                        r.font.size = Pt(11)
                        if i == 0:
                            r.bold = True
            doc.add_paragraph().paragraph_format.space_after = Pt(0)

    doc.save(destino)


# ------------------------------------------------------------------ html

ESTILO = """
@page { size: A4; margin: 3cm 2cm 2cm 3cm; }
html,body{background:#FFF}
body{font-family:"Times New Roman",Times,serif;font-size:12pt;line-height:1.5;
     text-align:justify;color:#000;margin:0}
h1{font-size:14pt;font-weight:700;text-transform:uppercase;text-align:left;
   margin:0 0 18pt;line-height:1.5}
h2{font-size:12pt;font-weight:700;margin:18pt 0 6pt;line-height:1.5}
h3{font-size:12pt;font-weight:700;font-style:italic;margin:18pt 0 6pt;line-height:1.5}
p{margin:0;text-indent:1.25cm}
blockquote{margin:6pt 0 0 2cm;font-style:italic}
blockquote p{text-indent:0}
ol,ul{margin:6pt 0 0 1.25cm;padding:0}
li{margin:0 0 6pt;line-height:1.5}
table{border-collapse:collapse;width:100%;margin:6pt 0;font-size:11pt}
th,td{border:1px solid #000;padding:4pt 6pt;text-align:left;line-height:1.2}
hr{display:none}
"""

def montar_html(blocos, destino: Path, titulo: str):
    def inline(t):
        t = re.sub(r"\*\*([^*]+)\*\*", r"<strong>\1</strong>", t)
        t = re.sub(r"\*([^*]+)\*", r"<em>\1</em>", t)
        t = re.sub(r"`([^`]+)`", r"\1", t)
        return t

    out, lista = [], None
    for tipo, conteudo in blocos:
        if tipo in ("num", "item"):
            tag = "ol" if tipo == "num" else "ul"
            if lista != tag:
                if lista: out.append(f"</{lista}>")
                out.append(f"<{tag}>"); lista = tag
            out.append(f"<li>{inline(conteudo)}</li>")
            continue
        if lista:
            out.append(f"</{lista}>"); lista = None
        if tipo == "h1": out.append(f"<h1>{inline(conteudo)}</h1>")
        elif tipo == "h2": out.append(f"<h2>{inline(conteudo)}</h2>")
        elif tipo == "h3": out.append(f"<h3>{inline(conteudo)}</h3>")
        elif tipo == "p": out.append(f"<p>{inline(conteudo)}</p>")
        elif tipo == "cita": out.append(f"<blockquote><p>{inline(conteudo)}</p></blockquote>")
        elif tipo == "tabela":
            linhas = [l for l in conteudo if not re.match(r"^\|[\s|:-]+\|$", l)]
            cels = [[c.strip() for c in l.strip("|").split("|")] for l in linhas]
            if not cels: continue
            out.append("<table>")
            for i, linha in enumerate(cels):
                tag = "th" if i == 0 else "td"
                out.append("<tr>" + "".join(f"<{tag}>{inline(c)}</{tag}>" for c in linha) + "</tr>")
            out.append("</table>")
    if lista: out.append(f"</{lista}>")

    io.open(destino, "w", encoding="utf-8").write(
        f"""<!doctype html><html lang="pt-BR"><head><meta charset="UTF-8">
<title>{titulo}</title><style>{ESTILO}</style></head><body>
{"".join(out)}
</body></html>""")


# ------------------------------------------------------------------ main

if __name__ == "__main__":
    entrada = Path(sys.argv[1])
    md = limpar(io.open(entrada, encoding="utf-8").read())
    blocos = partir(md)
    titulo = next((c for t, c in blocos if t == "h1"), entrada.stem)

    SAIDA.mkdir(exist_ok=True)
    base = SAIDA / entrada.stem
    montar_docx(blocos, base.with_suffix(".docx"), titulo)
    montar_html(blocos, base.with_suffix(".html"), titulo)

    palavras = len(re.findall(r"\w+", " ".join(c for t, c in blocos if isinstance(c, str))))
    print(f"{entrada.name}: {palavras} palavras -> {base.name}.docx e .html")
