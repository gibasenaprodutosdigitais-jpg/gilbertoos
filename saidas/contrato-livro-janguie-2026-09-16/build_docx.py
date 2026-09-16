# -*- coding: utf-8 -*-
"""
Contraproposta ao "Termo de Compromisso, Cessao de Direitos Autorais e
Responsabilidade para Participacao no Projeto Livro com Janguie".

Mantem a estrutura e a numeracao do documento original, reescrevendo as
clausulas criticas (marcadas [AJUSTADO]) para: preservar o credito e o
uso proprio do capitulo pelo coautor, eliminar o "cheque em branco" do
preco dos exemplares, condicionar a obrigacao de compra a inclusao
efetiva do capitulo, tornar a penalidade proporcional, e estreitar a
clausula de imagem. Nao e parecer juridico - e um ponto de partida pra
negociacao / revisao por advogado.
"""
import os
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, "Termo Coautoria - Contraproposta (revisado).docx")

TAG_COLOR = RGBColor(0xB0, 0x3A, 0x2E)   # vermelho queimado - marca o que mudou
INK = RGBColor(0x15, 0x15, 0x15)

doc = Document()

# ---------- estilo base ----------
normal = doc.styles["Normal"]
normal.font.name = "Times New Roman"
normal.font.size = Pt(12)
normal.paragraph_format.space_after = Pt(10)
normal.paragraph_format.line_spacing = 1.5

sec = doc.sections[0]
sec.left_margin = Cm(3); sec.top_margin = Cm(3)
sec.right_margin = Cm(2); sec.bottom_margin = Cm(2)


def h1(text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text)
    r.bold = True; r.font.size = Pt(13)
    return p


def h2(text, tag=False):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.bold = True; r.font.size = Pt(12)
    if tag:
        rt = p.add_run("  [AJUSTADO]")
        rt.bold = True; rt.font.size = Pt(10); rt.font.color.rgb = TAG_COLOR
    return p


def clause(num, text, tag=False, justify=True):
    p = doc.add_paragraph()
    if justify:
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.first_line_indent = Cm(1.25)
    r = p.add_run(f"{num}\t" if num else "")
    r.bold = bool(num)
    r2 = p.add_run(text)
    if tag:
        rt = p.add_run("  [AJUSTADO]")
        rt.bold = True; rt.font.size = Pt(9); rt.font.color.rgb = TAG_COLOR
    return p


def bullet(text):
    p = doc.add_paragraph(style="List Bullet")
    p.add_run(text)
    return p


def field_row(table, label, placeholder):
    row = table.add_row().cells
    row[0].text = label
    row[0].paragraphs[0].runs[0].bold = True
    row[1].text = placeholder


def rule():
    p = doc.add_paragraph()
    pPr = p._p.get_or_add_pPr()
    pBdr = pPr.makeelement(qn('w:pBdr'), {})
    bottom = pPr.makeelement(qn('w:bottom'), {
        qn('w:val'): 'single', qn('w:sz'): '6', qn('w:space'): '1', qn('w:color'): '999999'
    })
    pBdr.append(bottom)
    pPr.append(pBdr)


# ============================================================ CAPA / RESUMO
h1("PROPOSTA DE AJUSTES AO TERMO DE COAUTORIA")
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run('Projeto "Livro com Janguiê" — Empreendedorismo de Impacto')
r.italic = True

doc.add_paragraph()
h2("O que muda em relação ao Termo original enviado, e por quê")
bullet('Capítulo II (Cessão de Direitos): a cessão passa de "gratuita, perpétua e exclusiva" '
       'para não exclusiva e limitada à publicação como parte deste Livro. O COAUTOR mantém '
       'o direito de usar, publicar e reaproveitar o próprio capítulo em qualquer outro '
       'contexto (redes, cursos, livros próprios), a qualquer tempo.')
bullet('Cláusula 2.3: o nome do COAUTOR passa a constar como crédito real de coautoria '
       '(capa, folha de rosto, sumário) — não mais como simples "referência de contribuição" '
       'sem valor jurídico — e o COAUTOR pode usar esse crédito na própria divulgação.')
bullet('Cláusula 2.4: a renúncia a qualquer pretensão presente ou futura, judicial ou '
       'extrajudicial, é substituída por uma cláusula direta de não pagamento de royalties '
       '— sem abrir mão do direito de recorrer à Justiça em caso de descumprimento do próprio '
       'Termo.')
bullet('Cláusula 4.2.1: fica expresso que a não entrega do capítulo, ou a não inclusão dele no '
       'Livro por motivo não atribuível ao COAUTOR, libera automaticamente o COAUTOR da '
       'obrigação de comprar os 100 exemplares (Capítulo V). No original isso não estava '
       'conectado — dava pra cobrar a compra mesmo com o capítulo de fora.')
bullet('Cláusula 5.1: o preço do exemplar deixa de ser "a ser informado oportunamente" e passa '
       'a ser um valor fixo, conhecido e aceito na assinatura — e a obrigação de compra só '
       'nasce depois que o capítulo for efetivamente incluído na versão final aprovada.')
bullet('Cláusulas 5.1.1 e 5.2: o prazo de pagamento sobe de 5 dias úteis pra 30 dias corridos, '
       'com 10 dias de regularização antes de qualquer cobrança; e a multa de 10% seco vira '
       'multa moratória de 2% + juros de 1% ao mês — padrão de mercado, não punitivo.')
bullet('Cláusula 3.5: se o capítulo for excluído por suspeita de plágio antes da publicação, o '
       'COAUTOR passa a ter direito à devolução de qualquer valor já pago pelos exemplares '
       '— o original negava reembolso mesmo nesse cenário.')
bullet('Capítulo VI (Imagem): a proibição deixa de valer pra "qualquer opinião, crítica ou '
       'insinuação" sobre o Autor e seus "projetos e empreendimentos" (bem amplo, sem prazo), '
       'e passa a valer só pra crítica infundada ou ofensa sobre o Autor e o Livro, por até 12 '
       'meses após a publicação — preservando opinião legítima e crítica técnica. A regra passa '
       'a valer nos dois sentidos.')
bullet('Terminologia: padronizado "COAUTOR" do início ao fim (o original alternava com '
       '"Mentorando" nas cláusulas de prazo e penalidade).')
bullet('Cláusula 8.1 (foro): sugestão opcional — manter São Paulo/SP como regra, mas permitir '
       'ao COAUTOR optar pelo foro do seu domicílio quando for ele quem processa. É o ponto '
       'mais fácil de a Editora recusar; inclua se quiser, mas não é o ajuste mais importante.')

doc.add_paragraph()
p = doc.add_paragraph()
r = p.add_run("Este documento é uma minuta de contraproposta pra negociação — não substitui "
              "a revisão de um advogado antes da assinatura final, especialmente pelo Termo "
              "já nascer como título executivo extrajudicial.")
r.italic = True; r.font.size = Pt(10)

doc.add_page_break()

# ============================================================ TERMO REVISADO
h1("TERMO DE COMPROMISSO, CESSÃO DE DIREITOS AUTORAIS E\nRESPONSABILIDADE PARA "
   'PARTICIPAÇÃO NO PROJETO "LIVRO COM JANGUIÊ"')
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("(versão revisada — proposta do COAUTOR)")
r.italic = True; r.font.size = Pt(10)

doc.add_paragraph('Pelo presente instrumento particular ("Termo"):')

table = doc.add_table(rows=0, cols=2)
table.style = "Table Grid"
table.alignment = WD_TABLE_ALIGNMENT.CENTER
for label, ph in [
    ("NOME:", "{{Nome completo}}"), ("EMAIL:", "{{E-mail}}"), ("TELEFONE:", "{{Telefone}}"),
    ("NACIONALIDADE:", "{{Nacionalidade}}"), ("ESTADO CIVIL:", "{{Estado Civil}}"),
    ("OCUPAÇÃO:", "{{Ocupação}}"), ("CPF/MF:", "{{CPF}}"), ("CARTEIRA DE IDENTIDADE (RG):", "{{RG}}"),
    ("ENDEREÇO:", "{{Logradouro}}"), ("NÚMERO:", "{{Número}}"), ("COMPLEMENTO:", "{{Complemento}}"),
    ("CEP:", "{{CEP}}"), ("BAIRRO:", "{{Bairro}}"), ("CIDADE:", "{{Cidade}}"), ("ESTADO:", "{{Estado}}"),
]:
    field_row(table, label, ph)

doc.add_paragraph()
doc.add_paragraph('Doravante denominado "COAUTOR", declara, de forma expressa, para os devidos '
                   'fins de direito, o que segue:')

# ---- Capítulo I
doc.add_paragraph()
h2("CAPÍTULO I. DO OBJETO")
clause("1.1.", 'O COAUTOR foi convidado a participar do projeto editorial denominado '
       '"EMPREENDEDORISMO DE IMPACTO: O poder da influência, da inovação e da criação de '
       'valor para escalar negócios" ("Livro"), de autoria principal do Sr. Dr. José Janguiê '
       'Bezerra Diniz ("Autor"), com publicação prevista por meio da editora Novo Século '
       '("Editora"), mediante a redação de um capítulo inédito sobre o tema.')
clause("1.2.", 'O COAUTOR declara estar ciente de que sua participação ocorre a convite do '
       'Autor, no âmbito do Instituto Êxito de Empreendedorismo, na condição de coautor de '
       'capítulo próprio, sendo o Sr. Dr. José Janguiê Bezerra Diniz o autor principal do Livro.')
clause("1.3.", 'A aceitação deste Termo implica plena ciência, anuência e concordância do '
       'COAUTOR com todas as cláusulas e obrigações aqui estabelecidas.')
clause("1.4.", 'O COAUTOR autoriza que o Autor possa promover, a seu critério, o registro do '
       'Capítulo, do Livro ou do presente Termo junto ao Escritório de Direitos Autorais '
       '(art. 19 da Lei nº 9.610/1998), em Cartório de Títulos e Documentos ou junto a outros '
       'órgãos especializados, sem prejuízo dos direitos do COAUTOR previstos neste Termo.')
clause("1.5.", 'O COAUTOR participa da presente obra na condição de autor de capítulo próprio, '
       'com os direitos e resguardos previstos neste Termo, sendo sua contribuição autoral '
       'reconhecida de forma expressa no Livro, nos termos do Capítulo II.', tag=True)

# ---- Capítulo II
h2("CAPÍTULO II. DA CESSÃO DE DIREITOS AUTORAIS")
clause("2.1.", 'O COAUTOR é o legítimo titular dos direitos autorais patrimoniais e morais '
       'sobre o conteúdo do capítulo a ser entregue ("Capítulo"), cedendo à Editora e ao Autor, '
       'em caráter não exclusivo, os direitos patrimoniais necessários à publicação, '
       'reprodução, edição, comercialização, divulgação e distribuição do Capítulo como parte '
       'integrante do Livro, em qualquer formato ou meio, físico ou digital, no Brasil e no '
       'exterior. A presente cessão não é perpétua nem exclusiva, permanecendo o COAUTOR livre '
       'para, a qualquer tempo e independentemente de nova autorização, utilizar, publicar, '
       'adaptar, reproduzir ou explorar o mesmo conteúdo, integral ou parcialmente, em obras, '
       'materiais, cursos, palestras ou canais próprios, ressalvada apenas a publicação de obra '
       'concorrente que reproduza integralmente o Livro.', tag=True)
clause("2.2.", 'A cessão de que trata o item 2.1 abrange os direitos de tradução, adaptação e '
       'inclusão em reedições futuras do Livro, sempre como parte do Livro, sem prejuízo da '
       'titularidade e da faculdade de uso próprio do Capítulo pelo COAUTOR prevista no item '
       '2.1.', tag=True)
clause("2.3.", 'O nome do COAUTOR constará como coautor do Capítulo correspondente, com '
       'crédito de autoria expresso na capa, na folha de rosto, no sumário e no próprio '
       'capítulo do Livro, o que constitui reconhecimento de sua condição de coautor e poderá '
       'ser livremente utilizado pelo COAUTOR para fins de divulgação de sua atuação '
       'profissional, respeitado o disposto no Capítulo VI.', tag=True)
clause("2.4.", 'O COAUTOR não fará jus a royalties, participação em receitas ou lucros '
       'decorrentes da exploração comercial do Livro pela Editora, ressalvados os direitos de '
       'titularidade, crédito e uso próprio do Capítulo assegurados nos itens 2.1 a 2.3, e sem '
       'prejuízo do direito de o COAUTOR recorrer às vias administrativa e judicial em caso de '
       'descumprimento de qualquer disposição deste Termo.', tag=True)

# ---- Capítulo III
h2("CAPÍTULO III. DA ORIGINALIDADE E RESPONSABILIDADE SOBRE O CONTEÚDO")
clause("3.1.", 'O COAUTOR declara, sob as penas da lei, que o conteúdo do capítulo a ser '
       'entregue será integralmente original, inédito, de sua exclusiva autoria e isento de '
       'plágio.')
clause("3.2.", 'É expressamente vedada a utilização de conteúdo total ou parcialmente copiado, '
       'mesmo que com adaptações, de terceiros ou de fontes não autorizadas, sob pena de '
       'responsabilidade civil, penal e contratual.')
clause("3.3.", 'O COAUTOR assume, com exclusividade, responsabilidade jurídica, administrativa '
       'e moral sobre o conteúdo de sua contribuição, isentando o Autor e a Editora de ônus '
       'decorrente de demandas de terceiros relativas à violação de direitos autorais, '
       'propriedade intelectual, direitos de personalidade, difamação ou outras infrações '
       'legais comprovadamente atribuíveis ao conteúdo do Capítulo.')
clause("3.4.", 'Caso seja constatado, a qualquer tempo, plágio ou violação de direitos de '
       'terceiros, o Autor poderá, mediante notificação prévia e por escrito ao COAUTOR, com '
       'prazo de 10 (dez) dias para manifestação: (i) excluir o Capítulo do Livro, mesmo após '
       'sua publicação; e (ii) exigir do COAUTOR o ressarcimento dos prejuízos diretos e '
       'comprovadamente decorrentes da violação, incluindo honorários advocatícios e custas '
       'processuais, limitado ao valor efetivamente apurado em decisão judicial transitada em '
       'julgado ou em acordo formal entre as partes.', tag=True)
clause("3.5.", 'Caso a infração seja identificada antes da publicação do Livro, o Autor poderá '
       'excluir o Capítulo, mediante comunicação prévia e por escrito ao COAUTOR com a '
       'exposição dos motivos. Nesta hipótese, a Editora restituirá integralmente ao COAUTOR, '
       'em até 15 (quinze) dias, qualquer valor já desembolsado nos termos do Capítulo V até a '
       'data da exclusão.', tag=True)

# ---- Capítulo IV
h2("CAPÍTULO IV. DA ENTREGA E FORMATAÇÃO DO CAPÍTULO")
clause("4.1.", 'O Capítulo deverá ser redigido com observância das normas da ABNT, em fonte '
       'Times New Roman, tamanho 12, espaçamento 1,5, margens de 3 cm (superior e esquerda) e '
       '2 cm (inferior e direita), alinhamento justificado, recuo de parágrafo de 1,25 cm e '
       'numeração de páginas no canto superior direito, e deverá conter, em média, 10 (dez) '
       'páginas.')
clause("4.2.", 'O COAUTOR assume o compromisso de entregar o Capítulo finalizado, no formato '
       'exigido, até a data-limite de 30 de setembro de 2026, exclusivamente para o e-mail: '
       'sergio.murilo@institutoexito.com.br.', tag=True)
clause("4.2.1.", 'O não envio do Capítulo dentro do prazo acima, ou a não inclusão do Capítulo '
       'no Livro por qualquer motivo não imputável ao COAUTOR, exime o Autor de qualquer '
       'obrigação de inclusão do conteúdo na obra e exime, automática e simultaneamente, o '
       'COAUTOR de qualquer obrigação prevista no Capítulo V deste Termo, sem que tal fato '
       'gere direito, indenização ou expectativa adicional a qualquer das partes.', tag=True)

# ---- Capítulo V
h2("CAPÍTULO V. DA AQUISIÇÃO DE EXEMPLARES")
clause("5.1.", 'Como contrapartida à publicação do Capítulo como parte do Livro, o COAUTOR se '
       'compromete a adquirir 100 (cem) exemplares do Livro, diretamente junto à Editora, pelo '
       'preço de custo de R$ {{Valor unitário do exemplar}} por exemplar, totalizando R$ '
       '{{Valor total}}, valor este conhecido e aceito pelo COAUTOR na data de assinatura deste '
       'Termo. Esta obrigação somente se torna exigível após a efetiva inclusão do Capítulo do '
       'COAUTOR na versão final do Livro aprovada para publicação.', tag=True)
clause("5.1.1.", 'O não cumprimento desta obrigação, uma vez exigível nos termos do item 5.1, '
       'ensejará a aplicação de penalidade pecuniária correspondente à diferença entre o '
       'número de exemplares efetivamente adquiridos e os 100 (cem) exemplares de aquisição '
       'obrigatória, multiplicada pelo valor unitário fixado no item 5.1, a ser paga '
       'diretamente à Editora no prazo de 30 (trinta) dias corridos após notificação por '
       'escrito, a qual deverá conceder ao COAUTOR prazo adicional de 10 (dez) dias para '
       'regularização voluntária antes de qualquer medida de cobrança.', tag=True)
clause("5.2.", 'O inadimplemento da obrigação prevista no item 5.1.1, transcorrido o prazo de '
       'regularização, ensejará multa moratória de 2% (dois por cento) sobre o valor devido, '
       'acrescida de juros de mora de 1% (um por cento) ao mês e correção monetária pelo IPCA, '
       'sem prejuízo de despesas de cobrança comprovadamente incorridas.', tag=True)

# ---- Capítulo VI
h2("CAPÍTULO VI. DA PRESERVAÇÃO DA IMAGEM")
clause("6.1.", 'O COAUTOR compromete-se a não praticar, divulgar ou compartilhar, por qualquer '
       'meio, crítica infundada, ofensa ou insinuação difamatória dirigida especificamente ao '
       'Autor ou ao Livro, durante a vigência da parceria editorial de que trata este Termo e '
       'por até 12 (doze) meses após a publicação do Livro, ressalvadas manifestações de '
       'opinião legítima, crítica jornalística, técnica ou de mercado, e o exercício regular '
       'do direito de resposta.', tag=True)
clause("6.2.", 'É vedado ao COAUTOR utilizar o nome, imagem, logomarca, marca registrada ou '
       'qualquer referência ao Autor para fins promocionais ou comerciais, sem prévia e '
       'expressa autorização por escrito.')
clause("6.3.", 'O disposto neste Capítulo aplica-se reciprocamente ao Autor e à Editora em '
       'relação à imagem, à reputação e às atividades profissionais do COAUTOR.', tag=True)

# ---- Capítulo VII
h2("CAPÍTULO VII. DA NATUREZA JURÍDICA E EXECUTIVIDADE")
clause("7.1.", 'Este Termo obriga o COAUTOR nos termos aqui previstos, sendo certo que não '
       'estabelece qualquer vínculo empregatício, associativo, societário ou de representação '
       'entre o COAUTOR e o Autor.')
clause("7.2.", 'O presente Termo constitui título executivo extrajudicial, nos termos do '
       'artigo 784, III, do Código de Processo Civil, exigível em caso de descumprimento de '
       'qualquer de suas cláusulas, desde que precedido de notificação prévia e por escrito à '
       'parte inadimplente, com prazo mínimo de 10 (dez) dias para regularização voluntária, '
       'exceto nas hipóteses em que este Termo já preveja prazo específico.', tag=True)

# ---- Capítulo VIII
h2("CAPÍTULO VIII. DO FORO")
clause("8.1.", 'Para dirimir dúvidas, conflitos ou litígios oriundos deste instrumento, as '
       'partes elegem o foro da Comarca de São Paulo/SP, ressalvada ao COAUTOR a faculdade de '
       'optar pelo foro de seu domicílio nas ações em que figurar como autor.',
       tag=True)

doc.add_paragraph()
doc.add_paragraph('E, por estar ciente e de pleno acordo com todas as disposições, firmo o '
                   'presente Termo na presença de duas testemunhas abaixo assinadas.')
doc.add_paragraph()
p = doc.add_paragraph("São Paulo, {{Data}}.")
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
doc.add_paragraph()
doc.add_paragraph()
p = doc.add_paragraph("_______________________________________")
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p = doc.add_paragraph("{{Nome completo}} — COAUTOR")
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
doc.add_paragraph()
doc.add_paragraph("Testemunhas:")
doc.add_paragraph()
p = doc.add_paragraph("1. ______________________________\t\t2. ______________________________")

doc.save(OUT)
print("OK ->", OUT)
