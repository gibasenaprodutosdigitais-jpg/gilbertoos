# -*- coding: utf-8 -*-
"""
Slides de PALCO para a Imersão Reengenharia Empresarial — Bloco 1, Tópico 1.
Fonte: "evento Gilberto para validação dia 27-08.pdf" (Mapa Executivo).

Regra de design (deck de palco, nao carrossel/estudo):
  - poucas palavras por tela, letra enorme — pontuacao visual do que ele fala
  - fundo grafite constante, um acento dourado, leao minusculo no canto
  - Fraunces pras falas/perguntas de impacto, Inter pros rotulos pequenos
"""
import os
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

HERE = os.path.dirname(os.path.abspath(__file__))
LEAO = os.path.join(HERE, "assets", "leao.png")
OUT = os.path.join(HERE, "Topico 01 - O Sofisma do Faturamento.pptx")

DARK   = RGBColor(0x0F, 0x14, 0x19)
GOLD   = RGBColor(0xC9, 0xA2, 0x27)
PAPER  = RGBColor(0xF4, 0xF1, 0xEA)
MUTE   = RGBColor(0x8a, 0x83, 0x74)
SERIF  = "Fraunces"
SANS   = "Inter"

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]
SW, SH = prs.slide_width, prs.slide_height

LEAO_RATIO = 698 / 769  # h/w do arquivo assets/leao.png (mark)


def bg(slide, color=DARK):
    r = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SW, SH)
    r.fill.solid(); r.fill.fore_color.rgb = color
    r.line.fill.background(); r.shadow.inherit = False
    slide.shapes._spTree.remove(r._element)
    slide.shapes._spTree.insert(2, r._element)
    return r


def leao_mark(slide):
    w = Inches(0.62)
    h = Emu(int(w * LEAO_RATIO))
    slide.shapes.add_picture(LEAO, SW - w - Inches(0.5), SH - h - Inches(0.42), w, h)


def rule(slide, x, y, w=Inches(1.6)):
    r = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, Pt(3))
    r.fill.solid(); r.fill.fore_color.rgb = GOLD
    r.line.fill.background(); r.shadow.inherit = False
    return r


def box(slide, x, y, w, h, anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.LEFT):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    return tb, tf


def para(tf, first, text, size, color, font=SERIF, bold=True, italic=False,
         align=PP_ALIGN.LEFT, space_before=0, line_spacing=1.08, spacing=None):
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.alignment = align
    p.line_spacing = line_spacing
    if space_before:
        p.space_before = Pt(space_before)
    r = p.add_run(); r.text = text
    r.font.name = font; r.font.size = Pt(size); r.font.bold = bold
    r.font.italic = italic; r.font.color.rgb = color
    if spacing is not None:
        r._r.get_or_add_rPr().set('spc', str(int(spacing * 100)))
    return p


# ------------------------------------------------------------- SLIDE 1 · marcador do topico
s = prs.slides.add_slide(BLANK)
bg(s)
leao_mark(s)
_, tf = box(s, Inches(1.0), Inches(0.9), Inches(11.3), Inches(1.0), MSO_ANCHOR.TOP)
para(tf, True, "BLOCO 1 · 09H00 – 09H30", 16, GOLD, SANS, bold=True, spacing=3.0)

_, tf = box(s, Inches(1.0), Inches(2.5), Inches(11.3), Inches(3.4), MSO_ANCHOR.TOP)
para(tf, True, "TÓPICO 1", 22, MUTE, SANS, bold=True, spacing=4.0)
para(tf, False, "O Sofisma do\nFaturamento", 76, PAPER, SERIF, bold=True,
     space_before=14, line_spacing=1.02)

_, tf = box(s, Inches(1.02), Inches(5.75), Inches(3.2), Inches(0.1))
rule(s, Inches(1.02), Inches(5.78))

_, tf = box(s, Inches(1.0), Inches(5.95), Inches(10.5), Inches(0.8), MSO_ANCHOR.TOP)
para(tf, True, "A ilusão da empresa forte", 26, MUTE, SERIF, bold=False, italic=True)


# ------------------------------------------------------------- SLIDE 2 · a crenca a derrubar
s = prs.slides.add_slide(BLANK)
bg(s)
leao_mark(s)
_, tf = box(s, Inches(1.0), Inches(0.85), Inches(11.3), Inches(0.7), MSO_ANCHOR.TOP)
para(tf, True, "O QUE VOCÊ PENSA HOJE", 16, GOLD, SANS, bold=True, spacing=3.0)

_, tf = box(s, Inches(1.0), Inches(1.9), Inches(11.3), Inches(4.6), MSO_ANCHOR.MIDDLE)
para(tf, True,
     "“Minha empresa fatura\nR$ 500 mil por mês.\nEntão eu sei gerir —\ne estou protegido.”",
     54, PAPER, SERIF, bold=True, italic=True, line_spacing=1.12)


# ------------------------------------------------------------- SLIDE 3 · o confronto
s = prs.slides.add_slide(BLANK)
bg(s)
leao_mark(s)
_, tf = box(s, Inches(1.0), Inches(0.85), Inches(11.3), Inches(0.7), MSO_ANCHOR.TOP)
para(tf, True, "A VERDADE", 16, GOLD, SANS, bold=True, spacing=3.0)

_, tf = box(s, Inches(1.0), Inches(1.9), Inches(11.3), Inches(4.6), MSO_ANCHOR.MIDDLE)
para(tf, True, "Faturamento atrai\nconcorrente.", 68, PAPER, SERIF, bold=True, line_spacing=1.05)
para(tf, False, "Faturamento atrai\no Fisco.", 68, GOLD, SERIF, bold=True,
     space_before=28, line_spacing=1.05)


# ------------------------------------------------------------- SLIDE 4 · a escolha
s = prs.slides.add_slide(BLANK)
bg(s)
leao_mark(s)
_, tf = box(s, Inches(1.0), Inches(0.7), Inches(11.3), Inches(0.7), MSO_ANCHOR.TOP)
para(tf, True, "A ESCOLHA", 16, GOLD, SANS, bold=True, spacing=3.0)

_, tf = box(s, Inches(1.0), Inches(1.55), Inches(11.3), Inches(3.6), MSO_ANCHOR.TOP)
para(tf, True, "R$ 10 milhões faturados.", 44, PAPER, SERIF, bold=True, line_spacing=1.1)
para(tf, False, "R$ 800 mil sumindo em\nimposto errado e passivo.", 44, PAPER, SERIF,
     bold=True, space_before=10, line_spacing=1.1)
para(tf, False, "Ou uma estrutura enxuta e blindada.", 44, PAPER, SERIF, bold=True,
     space_before=26, line_spacing=1.1)

_, tf = box(s, Inches(1.0), Inches(6.0), Inches(11.3), Inches(1.0), MSO_ANCHOR.TOP)
para(tf, True, "VOCÊ ESCOLHE.", 34, GOLD, SANS, bold=True, spacing=1.5)


# ------------------------------------------------------------- SLIDE 5 · fechamento do topico
s = prs.slides.add_slide(BLANK)
bg(s)
leao_mark(s)
_, tf = box(s, Inches(1.0), Inches(0.85), Inches(11.3), Inches(0.7), MSO_ANCHOR.TOP)
para(tf, True, "PRA LEMBRAR", 16, GOLD, SANS, bold=True, spacing=3.0)

_, tf = box(s, Inches(1.0), Inches(2.0), Inches(11.3), Inches(4.2), MSO_ANCHOR.TOP)
para(tf, True, "Faturamento sem\nReengenharia é\napenas um\nrisco maior.", 64, PAPER, SERIF,
     bold=True, line_spacing=1.08)
rule(s, Inches(1.02), Inches(6.55))


prs.save(OUT)
print("OK ->", OUT)
print("slides:", len(prs.slides._sldIdLst))
